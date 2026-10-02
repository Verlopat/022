"""Audit a .docx for the black-and-white formatting rules."""
import re
import sys
import zipfile
from collections import Counter

from docx import Document
from docx.oxml.ns import qn

path = sys.argv[1]
z = zipfile.ZipFile(path)
xml = z.read('word/document.xml').decode('utf8')
print('parts with header/footer:', [n for n in z.namelist() if 'header' in n or 'footer' in n])
print('headerReference:', xml.count('headerReference'), 'footerReference:', xml.count('footerReference'),
      'pgNumType:', xml.count('pgNumType'), 'fldSimple/PAGE:', len(re.findall(r'PAGE', xml)))
print('underline w:u:', len(re.findall(r'<w:u [^>]*w:val="(?!none)', xml)))
cols = Counter(re.findall(r'<w:color w:val="([0-9A-Fa-f]{6}|auto)"', xml))
print('run colours:', dict(cols))
print('highlight:', xml.count('<w:highlight'), 'shading:', len(re.findall(r'<w:shd [^>]*w:fill="(?!FFFFFF|auto|ffffff)', xml)))
print('page borders:', xml.count('pgBorders'))
st = z.read('word/styles.xml').decode('utf8')
scol = Counter(re.findall(r'<w:color w:val="([0-9A-Fa-f]{6}|auto)"', st))
print('style colours:', dict(scol), 'style underline:', len(re.findall(r'<w:u [^>]*w:val="(?!none)', st)),
      'theme colours in styles:', st.count('themeColor'))
d = Document(path)
bad = 0
for t in d.tables:
    for i, row in enumerate(t.rows):
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    if i > 0 and r.bold and r.text.strip() and not r.text.strip().startswith('Algorithm'):
                        bad += 1
print('tables:', len(d.tables), 'bold runs in table bodies:', bad)
imgs = [n for n in z.namelist() if n.startswith('word/media/')]
print('images:', len(imgs))
heads = [p.text for p in d.paragraphs if p.style.name.startswith('Heading')]
print('headings:', len(heads))
nonbold = [p.text for p in d.paragraphs if p.style.name.startswith('Heading') and not all(r.bold for r in p.runs if r.text.strip())]
print('non-bold headings:', nonbold[:5])
